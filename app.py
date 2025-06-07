import streamlit as st
import re
import io
from docx import Document
from PyPDF2 import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFilter
import pytesseract
import cv2
import numpy as np
from moviepy.editor import ImageSequenceClip

# Extract text from Word documents
def extract_text_from_word(doc_file):
    doc = Document(doc_file)
    return [p.text for p in doc.paragraphs], doc

# Extract text from PDF documents
def extract_text_from_pdf(pdf_file):
    reader = PdfReader(pdf_file)
    return [page.extract_text() for page in reader.pages]

# Mask sensitive data in text
def mask_sensitive_data(paragraphs, selected_options):
    regex_patterns = {
        "Emails": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\\.com",
        "Phone Numbers": r"\\b\\d{10}\\b",
        "Account Numbers": r"\\b\\d{11,14}\\b",
        "IFSC Codes": r"\\b[A-Z]{4}\\d{7}\\b"
    }
    masked = []
    for para in paragraphs:
        for opt in selected_options:
            para = re.sub(regex_patterns.get(opt, ''), '[DATA HIDDEN]', para)
        masked.append(para)
    return masked

# Create Word doc from masked text
def create_word(masked_paragraphs):
    buffer = io.BytesIO()
    doc = Document()
    for para in masked_paragraphs:
        doc.add_paragraph(para)
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Create PDF from masked text
def create_pdf(masked_paragraphs):
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    text = c.beginText(40, 750)
    for para in masked_paragraphs:
        for line in para.split("\n"):
            text.textLine(line)
    c.drawText(text)
    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer

# Blur an area in an image
def blur_area(image, box):
    region = image.crop(box)
    blurred = region.filter(ImageFilter.GaussianBlur(radius=5))
    image.paste(blurred, box)

# Process single image for sensitive data
def process_image(image, selected_options):
    pytesseract.pytesseract.tesseract_cmd = r"C:\\Program Files\\Tesseract-OCR\\tesseract.exe"
    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
    patterns = {
        "Emails": r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\\.[a-zA-Z0-9-.]+",
        "Phone Numbers": r"\\b\\d{10}\\b",
        "Account Numbers": r"\\b\\d{11,14}\\b",
        "IFSC Codes": r"\\b[A-Z]{4}\\d{7}\\b"
    }
    for i, word in enumerate(data['text']):
        for opt in selected_options:
            if re.match(patterns.get(opt, ''), word):
                x, y, w, h = data['left'][i], data['top'][i], data['width'][i], data['height'][i]
                blur_area(image, (x, y, x+w, y+h))
    return image

# Process video frame-by-frame
def process_frame(frame, selected_options):
    image = Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
    return cv2.cvtColor(np.array(process_image(image, selected_options)), cv2.COLOR_RGB2BGR)

def process_video(video_path, selected_options):
    cap = cv2.VideoCapture(video_path)
    frames = []
    while True:
        ret, frame = cap.read()
        if not ret:
            break
        frames.append(process_frame(frame, selected_options))
    cap.release()
    return frames

st.title("Sensitive Data Masking in Files, Images, and Videos")

uploaded_file = st.file_uploader("Upload a Word, PDF, Image, or Video file", type=["docx", "pdf", "png", "jpg", "jpeg", "mp4", "avi", "mov"])

if uploaded_file:
    ftype = uploaded_file.type
    options = ["Emails", "Phone Numbers", "Account Numbers", "IFSC Codes"]

    if "wordprocessingml" in ftype:
        paras, _ = extract_text_from_word(uploaded_file)
    elif ftype == "application/pdf":
        paras = extract_text_from_pdf(uploaded_file)
    elif "image" in ftype:
        image = Image.open(uploaded_file)
        st.image(image, caption="Original Image", use_column_width=True)
        selected = st.multiselect("Select data types to blur in the image:", options)
        if st.button("Blur Sensitive Data"):
            result = process_image(image, selected)
            st.image(result, caption="Blurred Image", use_column_width=True)
            buf = io.BytesIO()
            result.save(buf, format="PNG")
            st.download_button("Download", buf, file_name="blurred.png")
        st.stop()
    elif "video" in ftype:
        st.video(uploaded_file)
        vpath = "temp_video.mp4"
        with open(vpath, "wb") as f:
            f.write(uploaded_file.getbuffer())
        selected = st.multiselect("Select data types to mask in video:", options)
        if st.button("Mask Video"):
            frames = process_video(vpath, selected)
            clip = ImageSequenceClip(frames, fps=24)
            out = "masked_video.mp4"
            clip.write_videofile(out)
            with open(out, "rb") as v:
                st.download_button("Download Masked Video", v, file_name=out)
        st.stop()
    else:
        st.error("Unsupported file type.")
        st.stop()

    st.subheader("Extracted Text")
    st.write("\n".join(paras))
    selected = st.multiselect("Select data types to mask:", options)
    if st.button("Mask Sensitive Data"):
        masked = mask_sensitive_data(paras, selected)
        st.subheader("Masked Text")
        st.write("\n".join(masked))
        if "wordprocessingml" in ftype:
            st.download_button("Download Word", create_word(masked), "masked.docx")
        elif ftype == "application/pdf":
            st.download_button("Download PDF", create_pdf(masked), "masked.pdf")
